import os

import pandas as pd
import xlwt
from openpyxl import load_workbook

from excel_splitter.core import split_workbook

legacy = os.path.join(os.getcwd(), 'sample_input_legacy.xls')
out = os.path.join(os.getcwd(), 'sample_output_legacy.xls')

legacy_df = pd.DataFrame([
    ['ACCOUNTING OFFICE', 'Employee A', 'PERMANENT'],
    ['ACCOUNTING OFFICE', 'Employee B', 'PERMANENT'],
    ['ACCOUNTING OFFICE', 'Employee C', 'UNKNOWN'],
    ['HR OFFICE', 'Employee D', 'PERMANENT'],
    ['HR OFFICE', 'Employee E', 'CONTRACTUAL'],
    ['ASSESSOR OFFICE', 'Employee F', 'PERMANENT'],
], columns=['Department', 'Name', 'ID Number'])

book = xlwt.Workbook()
sheet = book.add_sheet('Sheet1')
for col, header in enumerate(legacy_df.columns):
    sheet.write(0, col, header)
for row_num, row in enumerate(legacy_df.itertuples(index=False, name=None), start=1):
    for col_num, value in enumerate(row):
        sheet.write(row_num, col_num, value)
book.save(legacy)

result = split_workbook(legacy, 'Sheet1', 'Department', 'ID Number', out)
print('result', result)
print('output exists', os.path.exists(out))
print('sheet names', sorted([s.title for s in load_workbook(out).worksheets]))
