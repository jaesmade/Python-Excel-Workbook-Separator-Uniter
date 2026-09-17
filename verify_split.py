import os
from openpyxl import Workbook, load_workbook
from excel_splitter.core import split_workbook

source = os.path.join(os.getcwd(), 'sample_input.xlsx')
out = os.path.join(os.getcwd(), 'sample_output.xlsx')

wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'
ws.append(['Department', 'Name', 'ID Number'])
ws.append(['ACCOUNTING OFFICE', 'Employee A', 'PERMANENT'])
ws.append(['ACCOUNTING OFFICE', 'Employee B', 'PERMANENT'])
ws.append(['ACCOUNTING OFFICE', 'Employee C', 'UNKNOWN'])
ws.append(['HR OFFICE', 'Employee D', 'PERMANENT'])
ws.append(['HR OFFICE', 'Employee E', 'CONTRACTUAL'])
ws.append(['ASSESSOR OFFICE', 'Employee F', 'PERMANENT'])
wb.save(source)

result = split_workbook(source, 'Sheet1', 'Department', 'ID Number', out)
print('result', result)
wb2 = load_workbook(out)
print('worksheets', sorted([s.title for s in wb2.worksheets]))
wb2.close()
