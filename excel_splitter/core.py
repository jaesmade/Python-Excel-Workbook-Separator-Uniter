from __future__ import annotations

import os
from collections import OrderedDict
from typing import Callable

import pandas as pd
import xlwt
from openpyxl import Workbook

from .utils import detect_header_names, normalize_group_value, sanitize_sheet_name


class WorkbookProcessingError(ValueError):
    pass


ALL_SHEETS = "__ALL_SHEETS__"
SORT_ALPHABETICAL = "alphabetical"
SORT_COUNT_ASCENDING = "count_ascending"
SORT_COUNT_DESCENDING = "count_descending"


def sort_group_items(groups, sort_mode: str = SORT_ALPHABETICAL):
    items = list(groups.items())
    if sort_mode == SORT_COUNT_ASCENDING:
        return sorted(items, key=lambda item: len(item[1]))
    if sort_mode == SORT_COUNT_DESCENDING:
        return sorted(items, key=lambda item: len(item[1]), reverse=True)
    return sorted(items, key=lambda item: " - ".join(item[0]).casefold())


def _select_excel_engine(file_path: str) -> str | None:
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".xls":
        return "xlrd"
    if extension in {".xlsx", ".xlsm"}:
        return "openpyxl"
    return None


def _read_input_file(file_path: str, sheet_name: str | None = None) -> pd.DataFrame:
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".csv":
        return pd.read_csv(file_path)
    return pd.read_excel(file_path, sheet_name=sheet_name, engine=_select_excel_engine(file_path))


def _iter_input_tables(file_path: str, sheet_name: str | None):
    if os.path.splitext(file_path)[1].lower() == ".csv":
        yield "CSV file", _read_input_file(file_path)
        return

    engine = _select_excel_engine(file_path)
    with pd.ExcelFile(file_path, engine=engine) as workbook:
        selected_sheets = workbook.sheet_names if sheet_name == ALL_SHEETS else [sheet_name]
        for source_sheet_name in selected_sheets:
            yield source_sheet_name, pd.read_excel(file_path, sheet_name=source_sheet_name, engine=engine)


def _normalize_input_tables(input_tables):
    tables = []
    headers: list[str] = []
    seen_headers: set[str] = set()

    for source_name, table in input_tables:
        normalized_table = table.copy()
        normalized_headers = detect_header_names(normalized_table.columns)
        if len(normalized_headers) != len(set(normalized_headers)):
            raise WorkbookProcessingError(f"Duplicate column headers found in '{source_name}'.")

        normalized_table.columns = normalized_headers
        for header in normalized_headers:
            if header not in seen_headers:
                headers.append(header)
                seen_headers.add(header)
        tables.append((source_name, normalized_table))

    if not tables:
        return [], headers

    aligned_tables = [
        (source_name, table.reindex(columns=headers))
        for source_name, table in tables
    ]
    return aligned_tables, headers


def _ensure_python_value(value):
    if pd.isna(value):
        return None
    return value


def _iter_sheet_rows(sheet):
    if isinstance(sheet, pd.DataFrame):
        for row in sheet.itertuples(index=False, name=None):
            yield row
        return

    for row in sheet.iter_rows(values_only=True):
        yield row


def _first_row_headers(sheet) -> list[str]:
    if isinstance(sheet, pd.DataFrame):
        return detect_header_names(list(sheet.columns))

    rows = list(_iter_sheet_rows(sheet))
    if not rows:
        return []
    return detect_header_names(rows[0])


def read_workbook_sheet(file_path: str, sheet_name: str):
    data = _read_input_file(file_path, sheet_name)
    headers = _first_row_headers(data)
    if not headers or all((header == "" for header in headers)):
        raise WorkbookProcessingError("The selected worksheet does not contain any data.")
    return data, data, headers


def build_group_map(
    sheet,
    department_column: str,
    id_number_column: str | None = None,
    grouping_columns: list[str] | None = None,
):
    headers = _first_row_headers(sheet)
    if not headers:
        raise WorkbookProcessingError("The selected worksheet does not contain any data.")

    if grouping_columns is None:
        grouping_columns = [department_column]
        if id_number_column:
            grouping_columns.append(id_number_column)
    grouping_columns = [str(column).strip() for column in grouping_columns if str(column).strip()]
    if not grouping_columns:
        raise WorkbookProcessingError("At least one grouping column must be selected.")

    if isinstance(sheet, pd.DataFrame):
        columns = [str(col).strip() for col in list(sheet.columns)]
        missing = [column for column in grouping_columns if column not in columns]
        if missing:
            raise WorkbookProcessingError(f"Grouping column(s) not found in the worksheet headers: {', '.join(missing)}")
        grouping_indexes = [columns.index(column) for column in grouping_columns]
    else:
        missing = [column for column in grouping_columns if column not in headers]
        if missing:
            raise WorkbookProcessingError(f"Grouping column(s) not found in the worksheet headers: {', '.join(missing)}")
        grouping_indexes = [headers.index(column) for column in grouping_columns]

    groups = OrderedDict()

    for row in list(_iter_sheet_rows(sheet))[1:] if not isinstance(sheet, pd.DataFrame) else list(_iter_sheet_rows(sheet)):
        if not row or all((cell is None or pd.isna(cell) or str(cell).strip() == "") for cell in row):
            continue

        group_key = tuple(
            normalize_group_value(
                _ensure_python_value(row[index]) if index < len(row) else None,
                "UNKNOWN GROUP",
            )
            for index in grouping_indexes
        )

        if group_key not in groups:
            groups[group_key] = []
        groups[group_key].append(row)

    if not groups:
        raise WorkbookProcessingError("No data rows were found in the selected worksheet.")

    return headers, groups


def split_workbook(file_path: str, sheet_name: str, department_column: str, id_number_column: str | None,
                   output_file: str, selected_groups: set[tuple[str, str]] | None = None,
                   progress_callback: Callable[[int, str], None] | None = None,
           combine_duplicates: bool = True,
           grouping_columns: list[str] | None = None,
           sort_mode: str = SORT_ALPHABETICAL):
    return split_workbooks([file_path], sheet_name, department_column, id_number_column, output_file,
               selected_groups, progress_callback, combine_duplicates, grouping_columns, sort_mode)


def split_workbooks(file_paths: list[str], sheet_name: str, department_column: str, id_number_column: str | None,
                    output_file: str, selected_groups: set[tuple[str, str]] | None = None,
                    progress_callback: Callable[[int, str], None] | None = None,
                    combine_duplicates: bool = True,
                    grouping_columns: list[str] | None = None,
                    sort_mode: str = SORT_ALPHABETICAL):
    if not file_paths:
        raise WorkbookProcessingError("No files were selected for processing.")

    combined_groups: OrderedDict[tuple[str, ...], list[tuple]] = OrderedDict()
    total_files = len(file_paths)
    input_tables = []

    for file_index, file_path in enumerate(file_paths, start=1):
        for source_sheet_name, source_df in _iter_input_tables(file_path, sheet_name):
            if not source_df.empty:
                source_label = f"{os.path.basename(file_path)} - {source_sheet_name}"
                input_tables.append((source_label, source_df))

        if progress_callback:
            progress_callback(int((file_index / total_files) * 40), f"Reading: {os.path.basename(file_path)}")

    normalized_tables, combined_headers = _normalize_input_tables(input_tables)
    for source_label, source_df in normalized_tables:
        _, groups = build_group_map(source_df, department_column, id_number_column, grouping_columns)
        for group_key, rows in groups.items():
            if combine_duplicates:
                combined_groups.setdefault(group_key, []).extend(rows)
            else:
                unique_key = (*group_key, source_label)
                combined_groups.setdefault(unique_key, []).extend(rows)

    if not combined_groups:
        raise WorkbookProcessingError("No valid rows were found across the selected files.")

    if selected_groups is None:
        selected_groups = set(combined_groups.keys())

    if not selected_groups:
        raise WorkbookProcessingError("No groups were selected to generate worksheets.")

    output_ext = os.path.splitext(output_file)[1].lower()
    used_names: set[str] = set()
    processed_rows = 0
    ordered_groups = sort_group_items(combined_groups, sort_mode)
    total_groups = len(ordered_groups)

    if output_ext == ".xls":
        workbook = xlwt.Workbook()
        for index, (group_key, rows) in enumerate(ordered_groups, start=1):
            if group_key not in selected_groups:
                continue

            display_name = " - ".join(group_key)
            sheet_title = sanitize_sheet_name(display_name, used_names)
            sheet = workbook.add_sheet(sheet_title[:31])
            for col_num, header in enumerate(combined_headers or []):
                sheet.write(0, col_num, header)
            for row_num, row_values in enumerate(rows, start=1):
                for col_num, value in enumerate(row_values):
                    if pd.isna(value):
                        sheet.write(row_num, col_num, "")
                    else:
                        sheet.write(row_num, col_num, value)
                processed_rows += 1

            if progress_callback:
                progress_callback(40 + int((index / total_groups) * 60), f"Creating: {display_name}")

            workbook.save(output_file)
        return {
            "created_sheets": len(selected_groups),
            "processed_rows": processed_rows,
            "output_file": output_file,
        }

    output_wb = Workbook()
    default_sheet = output_wb.active
    output_wb.remove(default_sheet)

    for index, (group_key, rows) in enumerate(ordered_groups, start=1):
        if group_key not in selected_groups:
            continue

        display_name = " - ".join(group_key)
        sheet_title = sanitize_sheet_name(display_name, used_names)
        target_sheet = output_wb.create_sheet(title=sheet_title)

        for col_num, header in enumerate(combined_headers or [], start=1):
            target_sheet.cell(row=1, column=col_num, value=header)

        for row_num, row_values in enumerate(rows, start=2):
            for col_num, value in enumerate(row_values, start=1):
                if pd.isna(value):
                    target_sheet.cell(row=row_num, column=col_num, value=None)
                else:
                    target_sheet.cell(row=row_num, column=col_num, value=value)
            processed_rows += 1

        if progress_callback:
            progress_callback(40 + int((index / total_groups) * 60), f"Creating: {display_name}")

        output_wb.save(output_file)
    return {
        "created_sheets": len(output_wb.worksheets),
        "processed_rows": processed_rows,
        "output_file": output_file,
    }
